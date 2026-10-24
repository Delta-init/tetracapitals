"""names-to-cs.ts JSON → an Excel workbook (the user, 2026-10-24): every attendee with their student and CS, and a sheet by CS.

   python3 src/scripts/names-to-xlsx.py <names.json> <out.xlsx> [title]   (needs openpyxl)
"""
import json
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, dst, title = sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "Attendance"
data = json.load(open(src))
rows = data["rows"]

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
TONE = {"Found": "E8F5E9", "Not found": "FDECEA"}
POSSIBLE = "FFF8E1"


def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical="center")
    ws.freeze_panes = "A2"


def widths(ws):
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, *(len(str(c.value or "")) + 2 for c in col)))


wb = Workbook()
ws = wb.active
ws.title = "Students and CS"
header(ws, ["#", "Name in Meet", "Match", "Student", "Student code", "CS", "Team", "Location", "Email", "Phone", "Note"])
n, last = 0, None
for r in rows:
    if r["name"] != last:
        n += 1
        last = r["name"]
    note = "; ".join(x for x in [
        "CS switched off" if r.get("cs_off") else "",
        "marked 'not a student'" if r.get("not_student") else "",
        "pick the right one" if str(r["match"]).startswith("Possible") else "",
    ] if x)
    ws.append([n, r["name"], r["match"], r.get("student", ""), r.get("code", ""), r.get("cs", ""), r.get("team", ""),
               (r.get("location") or "").title(), r.get("email", ""), r.get("phone", ""), note])
    fill = PatternFill("solid", fgColor=TONE.get(r["match"], POSSIBLE))
    for c in ws[ws.max_row]:
        c.fill = fill
widths(ws)

by = wb.create_sheet("By CS")
header(by, ["CS", "Team", "Students (found)", "Names"])
groups = OrderedDict()
for r in rows:
    if r["match"] != "Found":
        continue
    groups.setdefault((r.get("cs", ""), r.get("team", "")), []).append(f'{r.get("student", "")} ({r.get("code", "")})')
for (cs, team), names in sorted(groups.items(), key=lambda kv: (-len(kv[1]), kv[0][0])):
    by.append([cs, team, len(names), ", ".join(names)])
    by[by.max_row][3].alignment = Alignment(wrap_text=True, vertical="top")
widths(by)
by.column_dimensions["D"].width = 90

found = len({r["name"] for r in rows if r["match"] == "Found"})
possible = len({r["name"] for r in rows if str(r["match"]).startswith("Possible")})
summary = wb.create_sheet("Summary", 0)
for line in [[title], [], ["Names in the attendance", data["names"]], ["Found (one student)", found],
             ["Several possible students — pick one", possible], ["Not found", data["names"] - found - possible],
             [], ["Looked up", data["at"][:16].replace("T", " ") + " UTC"]]:
    summary.append(line)
summary["A1"].font = Font(bold=True, size=14)
summary.column_dimensions["A"].width = 40
summary.column_dimensions["B"].width = 22

wb.save(dst)
print(f"Saved {dst}: {found} found, {possible} possible, {data['names'] - found - possible} not found")
